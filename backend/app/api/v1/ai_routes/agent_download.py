"""Agent物料多格式下载（Markdown/Word/PDF）"""
import re
from typing import Optional
from pathlib import Path
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import FileResponse, Response

from app.core.deps import get_current_user
from app.models import User
from app.services.ai.agent_materials import get_material_content

router = APIRouter()

# --------------------------------------------------------------------------- #
# Markdown 解析辅助函数
# --------------------------------------------------------------------------- #

def parse_markdown_to_text(md_content: str) -> list[dict]:
    """解析Markdown为结构化段落列表（用于docx和pdf生成）"""
    paragraphs = []
    lines = md_content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        
        # 标题
        if line.startswith('# '):
            paragraphs.append({'type': 'heading1', 'text': line[2:].strip()})
        elif line.startswith('## '):
            paragraphs.append({'type': 'heading2', 'text': line[3:].strip()})
        elif line.startswith('### '):
            paragraphs.append({'type': 'heading3', 'text': line[4:].strip()})
        # 列表
        elif line.startswith('- ') or line.startswith('* '):
            paragraphs.append({'type': 'list', 'text': line[2:].strip()})
        # 代码块
        elif line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            paragraphs.append({'type': 'code', 'text': '\n'.join(code_lines)})
            i += 1
            continue
        # 普通段落
        else:
            paragraphs.append({'type': 'paragraph', 'text': line})
        i += 1
    return paragraphs

# --------------------------------------------------------------------------- #
# 下载端点
# --------------------------------------------------------------------------- #

@router.get("/agents/materials/{ref}/download")
async def agent_download_material(
    ref: str,
    project_id: Optional[str] = Query(None),
    format: str = Query("md", regex="^(md|docx|pdf)$"),
    _: User = Depends(get_current_user),
):
    """下载Agent物料，支持 Markdown/Word/PDF 三种格式"""
    pid = project_id or "_global"
    m = get_material_content(pid, ref)
    
    if not m:
        raise HTTPException(status_code=404, detail="物料不存在")
    
    content = m["content"]
    
    if format == "md":
        return Response(
            content=content.encode("utf-8"),
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{ref}.md"'}
        )
    
    elif format == "docx":
        try:
            from docx import Document
            from docx.shared import Pt, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            
            doc = Document()
            
            for para in parse_markdown_to_text(content):
                if para['type'].startswith('heading'):
                    level = int(para['type'][-1])
                    doc.add_heading(para['text'], level=level)
                elif para['type'] == 'list':
                    doc.add_paragraph(para['text'], style='List Bullet')
                elif para['type'] == 'code':
                    p = doc.add_paragraph(para['text'])
                    p.style = 'Code'
                else:
                    doc.add_paragraph(para['text'])
            
            # 保存到临时文件
            import tempfile
            tmp_path = Path(tempfile.mktemp(suffix=".docx"))
            doc.save(str(tmp_path))
            
            return FileResponse(
                path=str(tmp_path),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                filename=f"{ref}.docx",
                deletion=True
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="python-docx未安装")
    
    elif format == "pdf":
        try:
            from fpdf import FPDF
            
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            
            # 使用默认字体（支持中文需要额外配置）
            pdf.set_font("Helvetica", size=10)
            
            for para in parse_markdown_to_text(content):
                if para['type'].startswith('heading'):
                    pdf.set_font("Helvetica", style='B', size=14)
                    pdf.cell(0, 10, para['text'], new_x="LMARGIN", new_y="NEXT")
                else:
                    pdf.set_font("Helvetica", size=10)
                    # 自动换行
                    pdf.multi_cell(0, 6, para['text'])
            
            # 保存到临时文件
            import tempfile
            tmp_path = Path(tempfile.mktemp(suffix=".pdf"))
            pdf.output(str(tmp_path))
            
            return FileResponse(
                path=str(tmp_path),
                media_type="application/pdf",
                filename=f"{ref}.pdf",
                deletion=True
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="fpdf2未安装")

